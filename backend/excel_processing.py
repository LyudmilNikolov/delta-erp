import math
import os
import re
from typing import Any

import numpy as np
import pandas as pd


REQUIRED_COLUMNS = {"Дата", "Код", "Стока", "Група", "Партида", "Количество", "Мярка"}

PHRASE_TO_TYPE = {
    "бут овк метро": "бут овк метро",
    "месо за готвене": "мг",
    "пържола бут": "пб",
    "долен бут": "дб",
    "едър замр": "едър замр",
    "мг окр": "мг окр",
    "пб окр": "пб окр",
    "стек овк": "стек овк",
    "филе окр": "филе окр",
    "шницел овк": "шницел овк",
    "чили и лимон": "чили и лимон",
    "бургер": "бургер",
    "кайма": "кайма",
    "кайми": "кайма",
    "кашкавал": "кашкавал",
    "кебапче": "кебапче",
    "кебапчета": "кебапче",
    "късчета": "късчета",
    "кюфте": "кюфте",
    "кюфтета": "кюфте",
    "мляно": "мляно",
    "наденица": "наденица",
    "пиле": "пиле",
    "плешка": "плешка",
    "руло": "руло",
    "селско": "селско",
    "стек": "стек",
    "сувлаки": "сувлаки",
    "филе": "филе",
    "шиш": "шиш",
    "шницел": "шницел",
}

GROUNDED_MEAT_TERMS = (
    "мляно",
    "кайма",
    "кайми",
    "кюфте",
    "кюфтета",
    "кебапче",
    "кебапчета",
    "бургер",
    "наденица",
    "наденички",
    "селско",
)
SEASONED_TERMS = (
    "овк",
    "овкус",
    "окр",
    "окрех",
    "билки",
    "мед и горчица",
    "чили и лимон",
    "кашкавал",
)

SEASONED_OR_GROUND_TYPES = {
    "бут овк метро",
    "бургер",
    "горен билки",
    "кайма",
    "кашкавал",
    "кебапче",
    "кюфте",
    "мг окр",
    "мляно",
    "наденица",
    "пб окр",
    "селско",
    "спеър рибс",
    "средиземн",
    "стек билки",
    "стек мед",
    "стек овк",
    "сувлаки",
    "трюфел",
    "филе окр",
    "чили и лимон",
    "шиш",
    "шиш мед",
    "шницел овк",
}


class ProcessingError(ValueError):
    """Expected workbook error that is safe to return as a client-facing 400."""


def _clean_text(value: Any) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _normalize_identifier(value: Any) -> str:
    text = _clean_text(_json_value(value))
    return re.sub(r"\.0$", "", text)


def _normalize_code(value: Any) -> str:
    return _normalize_identifier(value)


def _normalize_batch(value: Any) -> str:
    batch = _normalize_identifier(value)
    if not batch or batch.lower() in {"nan", "nat", "none", "null"}:
        return "NA"
    return batch


def _excel_engine(path: str) -> str:
    return "xlrd" if os.path.splitext(path)[1].lower() == ".xls" else "openpyxl"


def _leading_zero_width(number_format: Any) -> int | None:
    format_text = _clean_text(number_format).split(";", maxsplit=1)[0]
    format_text = re.sub(r'"[^"]*"|\[[^\]]*\]|\\.', "", format_text)
    integer_format = format_text.split(".", maxsplit=1)[0]
    zero_runs = re.findall(r"(?<![0#?])0{2,}(?![0#?])", integer_format)
    return max((len(run) for run in zero_runs), default=None)


def _format_excel_identifier(value: Any, number_format: Any = None) -> str:
    width = _leading_zero_width(number_format)
    if (
        width is not None
        and isinstance(value, (int, float, np.integer, np.floating))
        and not isinstance(value, (bool, np.bool_))
        and not pd.isna(value)
        and float(value).is_integer()
    ):
        integer = int(value)
        if integer < 0:
            return f"-{abs(integer):0{width}d}"
        return f"{integer:0{width}d}"
    return _normalize_identifier(value)


def _xlsx_displayed_codes(
    source_path: str,
    sheet_name: str,
    header: int,
    fallback_values: list[Any],
) -> list[str]:
    from openpyxl import load_workbook

    workbook = load_workbook(source_path, data_only=True, read_only=True)
    try:
        worksheet = workbook[sheet_name]
        header_row = header + 1
        code_column = next(
            (
                cell.column
                for cell in worksheet[header_row]
                if _clean_text(cell.value) == "Код"
            ),
            None,
        )
        if code_column is None:
            return [_normalize_code(value) for value in fallback_values]

        displayed_codes = []
        for offset, fallback in enumerate(fallback_values, start=1):
            cell = worksheet.cell(row=header_row + offset, column=code_column)
            value = fallback if cell.value is None else cell.value
            displayed_codes.append(_format_excel_identifier(value, cell.number_format))
        return displayed_codes
    finally:
        workbook.close()


def _xls_displayed_codes(
    source_path: str,
    sheet_name: str,
    header: int,
    fallback_values: list[Any],
) -> list[str]:
    import xlrd

    workbook = xlrd.open_workbook(source_path, formatting_info=True)
    worksheet = workbook.sheet_by_name(sheet_name)
    code_column = next(
        (
            column
            for column in range(worksheet.ncols)
            if _clean_text(worksheet.cell_value(header, column)) == "Код"
        ),
        None,
    )
    if code_column is None:
        return [_normalize_code(value) for value in fallback_values]

    displayed_codes = []
    for offset, fallback in enumerate(fallback_values, start=1):
        row = header + offset
        if row >= worksheet.nrows:
            displayed_codes.append(_normalize_code(fallback))
            continue

        cell = worksheet.cell(row, code_column)
        value = fallback if cell.value in {None, ""} else cell.value
        xf = workbook.xf_list[cell.xf_index]
        excel_format = workbook.format_map.get(xf.format_key)
        number_format = excel_format.format_str if excel_format else None
        displayed_codes.append(_format_excel_identifier(value, number_format))
    return displayed_codes


def _displayed_codes(
    source_path: str,
    sheet_name: str,
    header: int,
    fallback_values: list[Any],
) -> list[str]:
    try:
        if _excel_engine(source_path) == "xlrd":
            return _xls_displayed_codes(
                source_path,
                sheet_name,
                header,
                fallback_values,
            )
        return _xlsx_displayed_codes(
            source_path,
            sheet_name,
            header,
            fallback_values,
        )
    except Exception:
        return [_normalize_code(value) for value in fallback_values]


def _read_sheet_with_detected_header(source_path: str, sheet_name: str) -> pd.DataFrame:
    engine = _excel_engine(source_path)
    for header in range(6):
        try:
            df = pd.read_excel(source_path, sheet_name=sheet_name, header=header, engine=engine)
        except Exception:
            continue
        columns = set(str(column).strip() for column in df.columns)
        if REQUIRED_COLUMNS.issubset(columns):
            df.columns = [_clean_text(column) for column in df.columns]
            df["Код"] = _displayed_codes(
                source_path,
                sheet_name,
                header,
                df["Код"].tolist(),
            )
            return df
    raise ProcessingError(
        f"Листът „{sheet_name}“ не съдържа задължителните колони от Microinvest."
    )


def read_microinvest_sheet(source_path: str, sheet_name: str | None = None) -> tuple[str, pd.DataFrame]:
    try:
        excel = pd.ExcelFile(source_path, engine=_excel_engine(source_path))
    except Exception as exc:
        raise ProcessingError("Excel файлът не може да бъде прочетен.") from exc

    if sheet_name and sheet_name not in excel.sheet_names:
        raise ProcessingError(f"Листът „{sheet_name}“ не е намерен.")

    candidate_sheets = [sheet_name] if sheet_name else []
    if not candidate_sheets and "Microinvest" in excel.sheet_names:
        candidate_sheets.append("Microinvest")
    candidate_sheets.extend(sheet for sheet in excel.sheet_names if sheet not in candidate_sheets)

    errors = []
    for candidate in candidate_sheets:
        try:
            return candidate, _read_sheet_with_detected_header(source_path, candidate)
        except ProcessingError as exc:
            errors.append(str(exc))
    raise ProcessingError(
        "Няма лист със задължителните колони от Microinvest. " + " ".join(errors)
    )


def parse_package_weight_kg(description: Any) -> float:
    desc = _clean_text(description).lower()
    if not desc:
        return np.nan

    package_total_match = re.search(
        r"(?<!\d)(0[\.,]\d{2,3})\s*/\s*\d+\s*бр",
        desc,
    )
    if package_total_match:
        return float(package_total_match.group(1).replace(",", "."))

    multipack_match = re.search(
        r"(?<![\d\.,])(\d+)\s*(?:бр\.?\s*)?(?:по|[xх×*])\s*"
        r"(\d+(?:[\.,]\d+)?)\s*(кг|гр|г)?",
        desc,
    )
    if multipack_match:
        count = float(multipack_match.group(1))
        item_weight = float(multipack_match.group(2).replace(",", "."))
        unit = multipack_match.group(3)
        if unit in {"гр", "г"} or (not unit and item_weight >= 10):
            item_weight /= 1000
        return count * item_weight

    kilo_match = re.search(r"(\d+(?:[\.,]\d+)?)\s*кг", desc)
    if kilo_match:
        return float(kilo_match.group(1).replace(",", "."))

    gram_match = re.search(r"(\d+(?:[\.,]\d+)?)\s*гр", desc)
    if gram_match:
        return float(gram_match.group(1).replace(",", ".")) / 1000

    kg_match = re.search(r"(?<!\d)(0[\.,]\d{2,3})(?!\d)", desc)
    if kg_match:
        return float(kg_match.group(1).replace(",", "."))

    return np.nan


def infer_product_type(description: Any) -> str:
    desc = _clean_text(description).lower()

    has_seasoning = any(term in desc for term in ("овк", "овкус", "билки"))
    has_tenderizing = any(term in desc for term in ("окр", "окрех"))
    has_honey = "мед" in desc and "горчица" in desc

    if "трюфел" in desc:
        return "трюфел"
    if "спеър рибс" in desc:
        return "спеър рибс"
    if "средиземномор" in desc:
        return "средиземн"
    if "пилешко филе" in desc:
        return "филе пилешко"
    if "подбедрица" in desc:
        return "подбедрица"
    if "месо за готвене" in desc and has_tenderizing:
        return "мг окр"
    if "пържола" in desc and has_tenderizing:
        return "пб окр"
    if "филе" in desc and has_tenderizing:
        return "филе окр"
    if "шницел" in desc and has_seasoning:
        return "шницел овк"
    if "шиш" in desc and has_honey:
        return "шиш мед"
    if "стек" in desc and "горен бут" in desc and "билки" in desc:
        return "горен билки"
    if "стек" in desc and "плешка" in desc:
        return "плешка"
    if "стек" in desc and has_honey:
        return "стек мед"
    if "стек" in desc and has_seasoning:
        return "стек билки"
    if "шиш" in desc:
        return "шиш"
    if "обезкостен" in desc and "бут" in desc and has_seasoning:
        return "стек билки"
    if "селск" in desc and any(term in desc for term in ("кюфте", "кюфтета")):
        return "селско"
    if "наденич" in desc or "наденица" in desc:
        return "наденица"
    if "долно бутче" in desc:
        return "дб"
    if "месо за готвене" in desc:
        return "мг"
    if "стек" in desc:
        return "стек"
    if "обезкостен" in desc and "бут" in desc:
        return "пб"

    for phrase, product_type in sorted(PHRASE_TO_TYPE.items(), key=lambda item: len(item[0]), reverse=True):
        if phrase in desc:
            return product_type
    return desc


def classify_group(group: Any) -> str:
    group_text = _clean_text(group).lower()
    if any(term in group_text for term in ("стоки", "амадори", "патешко")):
        return "стоки"
    return "произв"


def is_seasoned_or_ground(description: Any, product_type: Any = None) -> bool:
    desc = _clean_text(description).lower()
    normalized_type = _clean_text(product_type).lower()
    return (
        normalized_type in SEASONED_OR_GROUND_TYPES
        or any(term in desc for term in GROUNDED_MEAT_TERMS + SEASONED_TERMS)
    )


def _convert_weight(row: pd.Series, row_number: int, warnings: list[dict[str, Any]]) -> float:
    quantity = pd.to_numeric(row.get("Количество"), errors="coerce")
    unit = _clean_text(row.get("Мярка")).lower().replace(" ", "")
    description = row.get("Стока")

    if pd.isna(quantity):
        warnings.append({
            "row": row_number,
            "message": "Липсващо или невалидно количество.",
            "product_name": _clean_text(description),
        })
        return np.nan

    if unit in {"кг", "кг."}:
        return float(quantity)

    if unit in {"бр", "бр."}:
        package_weight_kg = parse_package_weight_kg(description)
        if pd.isna(package_weight_kg):
            warnings.append({
                "row": row_number,
                "message": "Теглото на опаковката не може да бъде разпознато за продукт в бройки.",
                "product_name": _clean_text(description),
            })
            return np.nan
        return float(quantity) * float(package_weight_kg)

    warnings.append({
        "row": row_number,
        "message": f"Неподдържана мерна единица „{_clean_text(row.get('Мярка'))}“.",
        "product_name": _clean_text(description),
    })
    return np.nan


def _json_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, pd.Timestamp):
        return None if pd.isna(value) else value.date().isoformat()
    if isinstance(value, float) and math.isnan(value):
        return None
    if pd.isna(value):
        return None
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    return value


def dataframe_to_records(df: pd.DataFrame) -> list[dict[str, Any]]:
    return [
        {column: _json_value(value) for column, value in row.items()}
        for row in df.to_dict(orient="records")
    ]


def _build_original_row(row: pd.Series) -> dict[str, Any]:
    return {str(column): _json_value(value) for column, value in row.items()}


def normalize_microinvest(df: pd.DataFrame) -> tuple[pd.DataFrame, list[dict[str, Any]]]:
    missing = sorted(REQUIRED_COLUMNS - set(df.columns))
    if missing:
        raise ProcessingError(f"Липсват задължителни колони: {', '.join(missing)}")

    df = df.dropna(how="all").copy()
    df = df[df["Дата"].astype(str).str.strip().ne("Дата")]
    df = df[df["Стока"].apply(_clean_text).ne("")]
    warnings: list[dict[str, Any]] = []

    normalized = pd.DataFrame()
    normalized["date"] = pd.to_datetime(df["Дата"], dayfirst=True, errors="coerce")
    normalized["code"] = df["Код"].apply(_normalize_code)
    normalized["product_name"] = df["Стока"].apply(_clean_text)
    normalized["group"] = df.get("Група", "").apply(_clean_text)
    normalized["stock_type"] = df.get("Група", "").apply(classify_group)
    normalized["batch"] = df.get("Партида", "").apply(_normalize_batch)
    normalized["quantity"] = pd.to_numeric(df["Количество"], errors="coerce")
    normalized["unit"] = df["Мярка"].apply(_clean_text)
    normalized["weight_kg"] = [
        _convert_weight(row, int(index) + 2, warnings)
        for index, row in df.iterrows()
    ]

    if "вид" in df.columns:
        source_type = df["вид"].apply(_clean_text)
        inferred_type = df["Стока"].apply(infer_product_type)
        normalized["product_type"] = source_type.where(source_type.ne(""), inferred_type)
    else:
        normalized["product_type"] = df["Стока"].apply(infer_product_type)

    normalized["is_seasoned_or_ground"] = [
        is_seasoned_or_ground(description, product_type)
        for description, product_type in zip(
            normalized["product_name"],
            normalized["product_type"],
        )
    ]

    optional_map = {
        "document_no": "Документ №",
        "partner": "Партньор",
        "partner_group": "Група.1",
        "site": "Обект",
        "user": "Потребител",
        "operation": "Операция",
    }
    for output_col, source_col in optional_map.items():
        if source_col in df.columns:
            normalized[output_col] = df[source_col].apply(_clean_text)

    for index, date_value in normalized["date"].items():
        if pd.isna(date_value):
            warnings.append({
                "row": int(index) + 2,
                "message": "Датата не може да бъде разпозната.",
                "product_name": normalized.at[index, "product_name"],
            })

    normalized["original"] = [_build_original_row(row) for _, row in df.iterrows()]
    return normalized, warnings


def build_sheet1(normalized: pd.DataFrame) -> pd.DataFrame:
    grouped = normalized.pivot_table(
        index=["date", "code", "product_name", "stock_type"],
        values="weight_kg",
        aggfunc="sum",
    ).reset_index()
    return grouped.rename(columns={
        "date": "Дата",
        "code": "Код",
        "product_name": "Стока",
        "stock_type": "вид",
        "weight_kg": "Total",
    }).sort_values(["Дата", "Код", "Стока"])


def build_pivot(df: pd.DataFrame, seasoned_only: bool) -> pd.DataFrame:
    source = df[df["stock_type"].eq("произв")]
    if seasoned_only:
        source = source[source["is_seasoned_or_ground"]]
    pivot = source.pivot_table(
        index=["batch", "product_type"],
        values="weight_kg",
        aggfunc="sum",
    ).reset_index()
    pivot = pivot.rename(columns={
        "batch": "Партида",
        "product_type": "вид",
        "weight_kg": "Total",
    })
    return pivot.sort_values(["Партида", "вид"])


def write_result_workbook(
    result_path: str,
    sheet1: pd.DataFrame,
    seasoned_pivot: pd.DataFrame,
    all_kg_pivot: pd.DataFrame,
    details: pd.DataFrame,
    warnings: list[dict[str, Any]],
) -> None:
    excel_details = details.drop(columns=["original"], errors="ignore").copy()
    excel_details["date"] = excel_details["date"].dt.date

    with pd.ExcelWriter(result_path, engine="openpyxl") as writer:
        sheet1.to_excel(writer, sheet_name="Sheet1", index=False)
        seasoned_pivot.to_excel(writer, sheet_name="Sheet2", index=False)
        all_kg_pivot.to_excel(writer, sheet_name="Sheet3", index=False)
        excel_details.to_excel(writer, sheet_name="Details", index=False)
        if warnings:
            pd.DataFrame(warnings).to_excel(writer, sheet_name="Warnings", index=False)


def process_source(
    source_path: str,
    sheet_name: str | None = None,
    result_path: str | None = None,
) -> dict[str, Any]:
    selected_sheet, raw_df = read_microinvest_sheet(source_path, sheet_name)
    details, warnings = normalize_microinvest(raw_df)
    sheet1 = build_sheet1(details)
    seasoned_pivot = build_pivot(details, seasoned_only=True)
    all_kg_pivot = build_pivot(details, seasoned_only=False)

    if result_path:
        write_result_workbook(result_path, sheet1, seasoned_pivot, all_kg_pivot, details, warnings)

    return {
        "sheet_name": selected_sheet,
        "details": details,
        "sheet1": sheet1,
        "seasoned_pivot": seasoned_pivot,
        "all_kg_pivot": all_kg_pivot,
        "warnings": warnings,
        "seasoned_terms": {
            "ground_meats": list(GROUNDED_MEAT_TERMS),
            "seasoned": list(SEASONED_TERMS),
        },
    }
